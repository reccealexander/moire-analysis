#!/usr/bin/env python3
"""
Append an "All images" sheet to moire_grid_analysis.xlsx.

One row per .ibw file examined, across both samples, recording acquisition
metadata, the fit result where one was obtained, and an explicit reason where
none was. Files that produced no fit are kept as rows rather than dropped, so
the sheet accounts for every image rather than only the successful ones.
"""
import glob
import os
import re

import numpy as np
from igor2 import binarywave as bw
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from skimage.feature import peak_local_max

from moire_prep import load_ibw, level, nlm
from moire_calc import analyze, MATERIALS

from paths import RAW, WORKBOOK

MOS2_DIR = RAW["MoS2_2L"]
WSE2_DIR = RAW["WSe2"]
BOOK = WORKBOOK

GRID = {5: (14, 17), 2: (18, 21), 3: (23, 27), 6: (28, 31), 9: (32, 36),
        8: (37, 40), 7: (41, 45), 4: (46, 51), 1: (52, 64)}
POSITION = {1: ("532", "-1.69"), 2: ("532", "-1.5"), 3: ("532", "-1.29"),
            4: ("732", "-1.69"), 5: ("732", "-1.5"), 6: ("732", "-1.29"),
            7: ("932", "-1.69"), 8: ("932", "-1.49"), 9: ("932", "-1.29")}
L_MIN, L_MAX, S_MAX = 13.0, 21.0, 1.5
# |k| search windows, 1/nm -- identical to those used for the reported analyses
BAND_MOS2 = (0.025, 0.12)
BAND_WSE2 = (0.15, 0.45)

HEAD_FILL = PatternFill("solid", fgColor="1F3864")
HEAD_FONT = Font(bold=True, color="FFFFFF", size=11)
NOFIT_FILL = PatternFill("solid", fgColor="EFEFEF")
EXCL_FILL = PatternFill("solid", fgColor="FCE4E4")
THIN = Side(style="thin", color="BBBBBB")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def grid_of(num):
    for g, (lo, hi) in GRID.items():
        if lo <= num <= hi:
            return g
    return None


def raw_meta(path):
    """Shape, scan size and channel labels without assuming it is a real scan."""
    wave = bw.load(path)["wave"]
    arr = np.array(wave["wData"])
    note = wave["note"]
    note = note.decode("ascii", "ignore") if isinstance(note, bytes) else str(note)
    labels = [b.decode("ascii", "ignore") for ch in wave["labels"] for b in ch if b]
    m = re.search(r"ScanSize[:=]\s*([\d.eE+-]+)", note)
    scan = float(m.group(1)) * 1e9 if m else None
    return arr.shape, scan, labels


def fit(path, a_lat, band):
    """Attempt a hexagon fit; return (period, twist, strain) or None.

    band is the (lo, hi) |k| search window in 1/nm. It must match the window
    used by the analysis that produced the other sheets, otherwise the fitter
    finds different peaks and the sheets disagree.
    """
    got = load_ibw(path, "LateralTrace")
    if got is None:
        return None
    raw, scan, _ = got
    n = raw.shape[0]
    px = scan * 1e9 / n
    den = nlm(level(raw), 1.15)
    win = np.outer(np.hanning(n), np.hanning(n))
    spec = np.abs(np.fft.fftshift(np.fft.fft2(den * win)))
    c = n // 2
    yy, xx = np.mgrid[0:n, 0:n]
    fx = (xx - c) / (n * px)
    fy = (yy - c) / (n * px)
    rad = np.hypot(fx, fy)
    lo, hi = band

    peaks = peak_local_max(np.where((rad > lo) & (rad < hi), spec, 0),
                           min_distance=3, num_peaks=25)
    cands = []
    for y, x in peaks:
        w = 2
        sub = spec[y - w:y + w + 1, x - w:x + w + 1]
        if sub.shape != (2 * w + 1, 2 * w + 1):
            continue
        gy, gx = np.mgrid[-w:w + 1, -w:w + 1]
        tot = sub.sum()
        cands.append((np.array([(x + (sub * gx).sum() / tot - c) / (n * px),
                                (y + (sub * gy).sum() / tot - c) / (n * px)]), spec[y, x]))

    def amp(v):
        X = int(round(v[0] * n * px + c))
        Y = int(round(v[1] * n * px + c))
        if 0 <= X < n and 0 <= Y < n:
            return spec[max(0, Y - 1):Y + 2, max(0, X - 1):X + 2].max()
        return 0.0

    best = None
    for i, (b1, a1) in enumerate(cands):
        for j, (b2, a2) in enumerate(cands):
            if i == j:
                continue
            if np.hypot(*(b1 - b2)) < 0.02 or np.hypot(*(b1 + b2)) < 0.02:
                continue
            b3 = -(b1 + b2)
            k = [np.hypot(*b1), np.hypot(*b2), np.hypot(*b3)]
            if min(k) < lo or max(k) > hi or max(k) / min(k) > 1.5:
                continue
            score = min(a1, a2, amp(b3))
            if best is None or score > best[0]:
                best = (score, b1, b2)
    if best is None:
        return None
    res = analyze(best[1], best[2], a_lat, "recip")
    if not res["solutions"]:
        return None
    sol = res["solutions"][0]
    return res["lmean"], sol["twist"], sol["hetero"] * 100


def collect():
    rows = []
    for path in sorted(glob.glob(os.path.join(MOS2_DIR, "*.ibw"))):
        base = os.path.basename(path)
        shape, scan, labels = raw_meta(path)
        num = int(re.search(r"(\d{4})\.ibw", base).group(1))
        if len(shape) != 3 or shape[2] < 5 or scan is None:
            rows.append(["2L MoS2", base, None, shape[0] if shape else None, None,
                         ",".join(labels)[:40], None, None, None, None, None, None,
                         "not a scan (derived/exported field)"])
            continue
        g = grid_of(num)
        rc, cc = POSITION[g] if g else (None, None)
        px = scan / shape[0]
        got = fit(path, MATERIALS["MoS2"], BAND_MOS2)
        if got is None:
            note = "no hexagon found"
            rows.append(["2L MoS2", base, round(scan, 1), shape[0], round(px, 3),
                         "LateralTrace", g, rc, cc, None, None, None, note])
            continue
        L, tw, he = got
        if g is None:
            used, note = "no", "overview scan, outside the 3x3 grid"
        elif L_MIN <= L <= L_MAX and he < S_MAX:
            used, note = "yes", ""
        else:
            used = "no"
            note = ("period far from the ~16-18 nm cluster"
                    if not (L_MIN <= L <= L_MAX) else
                    "strain-dominated root, disagrees with repeats at this position")
        rows.append(["2L MoS2", base, round(scan, 1), shape[0], round(px, 3),
                     "LateralTrace", g, rc, cc, round(L, 2), round(tw, 3),
                     round(he, 2), note if note else "included in grid medians"])
        rows[-1].insert(12, used)

    for path in sorted(glob.glob(os.path.join(WSE2_DIR, "*.ibw"))):
        base = os.path.basename(path)
        shape, scan, labels = raw_meta(path)
        px = scan / shape[0]
        got = fit(path, MATERIALS["WSe2"], BAND_WSE2)
        layer = re.search(r"_(\d)L", base)
        sample = f"WSe2 {layer.group(1)}L" if layer else "WSe2"
        if got is None:
            rows.append([sample, base, round(scan, 1), shape[0], round(px, 3),
                         "LateralTrace", None, None, None, None, None, None, "no",
                         "no hexagon found"])
            continue
        L, tw, he = got
        rows.append([sample, base, round(scan, 1), shape[0], round(px, 3),
                     "LateralTrace", None, None, None, round(L, 2), round(tw, 3),
                     round(he, 2), "no",
                     "separate sample; multi-moire stack, strain unreliable"])
    return rows


def main():
    rows = collect()
    wb = load_workbook(BOOK)
    if "All images" in wb.sheetnames:
        del wb["All images"]
    ws = wb.create_sheet("All images")

    ws["A1"] = "Every image examined"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = ("One row per .ibw file opened during the analysis, across both samples. "
                "Rows with no fit are kept and the reason given. Grey = no fit, "
                "pink = fitted but excluded from the grid medians.")
    ws["A2"].font = Font(italic=True, size=9, color="555555")
    ws["A3"] = ("All fits are from FFT peak positions. Twist uses the exact polar "
                "decomposition; a = 0.3160 nm for MoS2, 0.3282 nm for WSe2.")
    ws["A3"].font = Font(italic=True, size=9, color="555555")

    head = ["Sample", "File", "Scan size (nm)", "Pixels", "nm/px", "Channel",
            "Grid #", "Row coord", "Col coord", "Moire period (nm)", "Twist (deg)",
            "Heterostrain (%)", "Used in medians", "Note"]
    ws.append([])
    ws.append(head)
    for c in range(1, len(head) + 1):
        cell = ws.cell(row=ws.max_row, column=c)
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER

    first = ws.max_row + 1
    for r in rows:
        while len(r) < len(head):
            r.append("")
        ws.append(r[:len(head)])
        rr = ws.max_row
        nofit = r[9] in (None, "")
        excluded = (not nofit) and r[12] == "no"
        for c in range(1, len(head) + 1):
            cell = ws.cell(row=rr, column=c)
            cell.border = BORDER
            cell.alignment = Alignment(horizontal="center" if c != len(head) else "left")
            if nofit:
                cell.fill = NOFIT_FILL
            elif excluded:
                cell.fill = EXCL_FILL

    for i, w in enumerate([11, 26, 14, 8, 8, 13, 8, 11, 11, 17, 12, 16, 15, 52], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = f"A{first}"
    ws.auto_filter.ref = f"A{first-1}:{get_column_letter(len(head))}{ws.max_row}"

    wb.save(BOOK)
    fitted = sum(1 for r in rows if r[9] not in (None, ""))
    print(f"wrote 'All images' sheet: {len(rows)} rows, {fitted} with a fit")


if __name__ == "__main__":
    main()
